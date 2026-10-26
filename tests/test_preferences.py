import copy
from PySide6.QtCore import Qt,QSettings
from PySide6.QtWidgets import QApplication,QFileDialog,QDialog,QPushButton,QLabel
from PySide6.QtTest import QTest
import pytest
from adv_editor.gui import Window
from adv_editor.core import sample,new_row,Database,statement,export_rpy
from adv_editor.i18n import tr,set_language


@pytest.fixture
def window(tmp_path):
    w=Window();w.db=Database(tmp_path/'Moon.advdb');w.db.save(w.data);w.refresh();w.show()
    QApplication.processEvents()
    yield w
    w.commit();w.dirty=False;w.close();QApplication.processEvents()


def test_id_visibility_does_not_change_data_and_tracks_edits(window):
    w=window;before=copy.deepcopy(w.data)
    assert not w.show_ids.isChecked()
    assert ' id d_' not in w.table.item(0,0).text()
    w.show_ids.setChecked(True)
    assert ' id '+w.rows[0]['id'] in w.table.item(0,0).text()
    w.show_ids.setChecked(False)
    assert w.data==before and not w.dirty
    w.table.setCurrentCell(0,2);w.table.editItem(w.table.item(0,2))
    w.editor.setPlainText('new text')
    assert ' id d_' not in w.table.item(0,0).text()
    w.commit()
    assert ' id d_' not in w.table.item(0,0).text()
    assert w.rows[0]['id']==before['sections'][0]['rows'][0]['id']


def test_export_id_switch_independent_of_preview(window,tmp_path,monkeypatch):
    w=window
    w.export_ids.setChecked(False)
    w.show_ids.setChecked(True)
    monkeypatch.setattr(QFileDialog,'getExistingDirectory',lambda *a,**k:str(tmp_path))
    w.export()
    source=(tmp_path/'adv_generated/dialogue.rpy').read_text()
    assert ' id d_' not in source
    w.export_ids.setChecked(True);w.show_ids.setChecked(False);w.export()
    assert ' id '+w.rows[0]['id'] in (tmp_path/'adv_generated/dialogue.rpy').read_text()


def test_menu_prompt_id_export_switch(tmp_path):
    data=sample();menu=dict(new_row(),text='Choice?',kind='menu',choices=[dict(text='Go',target='greeting')])
    data['sections'][0]['rows'].append(menu)
    path=export_rpy(data,tmp_path,include_ids=False)
    text=(path/'dialogue.rpy').read_text()
    assert '"Choice?"\n    menu:' in text and ' id d_' not in text
    assert 'jump adv_greeting' in text
    assert statement(menu,include_ids=True).startswith('"Choice?" id '+menu['id'])


def test_language_changes_preserve_unsaved_project_and_selection(window):
    from adv_editor.characters import CharacterDialog
    from adv_editor.sections import SectionForm
    from adv_editor.stage import StageDialog
    w=window
    w.table.setCurrentCell(0,2);w.table.editItem(w.table.item(0,2))
    w.editor.setPlainText('未儲存的台詞 {literal}')
    saved_id=w.rows[0]['id']
    w.language_combo.setCurrentIndex(w.language_combo.findData('en'))
    assert w.rows[0]['id']==saved_id and '未儲存' in w.rows[0]['text']
    before=copy.deepcopy(w.data)
    assert w.section_button.text()=='Chapters / Sections'
    assert w.table.horizontalHeaderItem(2).text()=='Dialogue'
    assert 'Ren’Py Text Editor' in w.windowTitle()
    assert w.project_location.text()==str(w.db.path.parent)
    assert w.mode.itemText(0)=='Visual mode'
    assert w.sections.currentText().startswith('第一章 / 相遇')
    characters=CharacterDialog(w.data,w)
    assert characters.windowTitle()=='Characters'
    assert characters.table.item(0,1).text()=='秋'
    assert characters.table.horizontalHeaderItem(1).text()=='Display name'
    form=SectionForm(w.data,parent=w)
    assert form.windowTitle()=='Add chapter / section'
    stage=StageDialog(w.data,w.db.path,parent=w)
    assert stage.kind.itemText(stage.kind.findData('scene'))=='Change background'
    assert stage.target.itemData(1)=='greeting'
    characters.reject();form.reject();stage.reject()
    w.language_combo.setCurrentIndex(w.language_combo.findData('zh_CN'))
    assert w.section_button.text()=='章节／段落管理'
    assert w.table.horizontalHeaderItem(2).text()=='台词内容'
    assert '文本编辑器' in w.windowTitle()
    assert w.show_ids.text()=='显示台词 ID'
    assert w.mode.itemText(0)=='视觉模式'
    assert w.data==before
    w.language_combo.setCurrentIndex(w.language_combo.findData('zh_TW'))
    assert w.section_button.text()=='章節／段落管理'
    assert w.data==before


def test_language_persistence(window):
    window.language_combo.setCurrentIndex(window.language_combo.findData('en'))
    assert QSettings('ADVScript','Editor').value('language')=='en'
    other=Window()
    assert other.section_button.text()=='Chapters / Sections'
    other.dirty=False;other.close()


def test_document_export_dialog_title_validation(window):
    from adv_editor.export_options import DocumentExportDialog
    dialog=DocumentExportDialog('Moon',False,window)
    dialog.title_input.setText('   ');dialog.accept()
    assert dialog.result()!=QDialog.DialogCode.Accepted
    assert dialog.error.text()
    dialog.title_input.setText('月光劇本');dialog.accept()
    assert dialog.result()==QDialog.DialogCode.Accepted


def test_pdf_custom_title_and_hidden_ids(window,tmp_path,monkeypatch):
    from adv_editor.export_options import DocumentExportDialog
    from PySide6.QtPdf import QPdfDocument
    def choose(dialog):
        dialog.title_input.setText('Moonlight Script')
        dialog.include_ids.setChecked(False)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(DocumentExportDialog,'exec',choose)
    path=tmp_path/'moon.pdf'
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a,**k:(str(path),''))
    window.document_out('pdf')
    pdf=QPdfDocument();pdf.load(str(path))
    text='\n'.join(pdf.getAllText(page).text() for page in range(pdf.pageCount()))
    assert 'Moonlight Script' in text
    assert window.rows[0]['id'] not in text
    assert '你好' in text
    assert pdf.metaData(QPdfDocument.MetaDataField.Title)=='Moonlight Script'
    pdf.close()


def test_word_title_and_id_checkbox(tmp_path):
    from adv_editor.documents import export_word
    from docx import Document
    path=tmp_path/'script.docx';data=sample()
    export_word(data,path,title='新標題',include_ids=False)
    paragraphs=[p.text for p in Document(path).paragraphs]
    assert paragraphs[0]=='新標題'
    assert not any(row['id'] in '\n'.join(paragraphs) for section in data['sections'] for row in section['rows'])
