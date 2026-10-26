import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import copy
from pathlib import Path
import pytest
from PySide6.QtWidgets import QApplication,QFileDialog,QDialog
from PySide6.QtGui import QTextCursor,QImage,QColor
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from adv_editor.core import sample,new_row,Database,export_rpy,export_excel,import_excel,validate,statement
from adv_editor.assets import import_font,font_file,asset_folder,copy_assets,import_image,image_file
from adv_editor.gui import Window

from adv_editor.platform_fonts import importable_system_font
FONT=importable_system_font()


@pytest.fixture(scope='module')
def app(): return QApplication.instance() or QApplication([])


@pytest.fixture
def window(app,tmp_path):
    w=Window();w.db=Database(tmp_path/'project.advdb');w.db.save(w.data);w.show();app.processEvents()
    yield w
    w.commit();w.dirty=False;w.close();app.processEvents()


def editor(w):
    w.table.setCurrentCell(0,2);w.table.editItem(w.table.item(0,2));return w.editor


def test_click_wait_is_same_say_statement(window,tmp_path):
    e=editor(window);e.set_source('前半後半')
    cursor=e.textCursor();cursor.setPosition(2);e.setTextCursor(cursor)
    QTest.mouseClick(window.click_wait_button,Qt.MouseButton.LeftButton)
    assert e.source()=='前半{w}後半'
    assert '{w}' not in e.toPlainText()
    window.commit();window.save()
    folder=export_rpy(window.data,tmp_path)
    assert '"前半{w}後半" id '+window.rows[0]['id'] in (folder/'dialogue.rpy').read_text()


def test_import_font_save_as_export_backup(window,tmp_path,monkeypatch):
    assert FONT.is_file()
    source=tmp_path/'original.ttf';source.write_bytes(FONT.read_bytes())
    e=editor(window);e.set_source('Font test');e.selectAll()
    monkeypatch.setattr(QFileDialog,'getOpenFileName',lambda *a,**k:(str(source),''))
    window.import_project_font()
    filename=next(iter(window.data['fonts']))
    assert font_file(window.db.path,filename).read_bytes()==source.read_bytes()
    assert 'font=adv_generated/fonts/'+filename in e.source()
    assert '{font' not in e.toPlainText()
    source.unlink()  # The project no longer depends on the original location.
    destination=tmp_path/'elsewhere'/'renamed.advdb';destination.parent.mkdir()
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a,**k:(str(destination),''))
    assert window.save(True)
    assert font_file(destination,filename).is_file()
    output=export_rpy(window.data,tmp_path/'game',destination)
    assert (output/'fonts'/filename).is_file()
    window.db.backup()
    backup=sorted((destination.parent/'renamed_backups').glob('*.advdb'))[-1]
    assert font_file(backup,filename).is_file()
    xlsx=tmp_path/'manuscript.xlsx'
    export_excel(window.data,xlsx,destination)
    assert import_excel(xlsx)==window.data
    assert font_file(xlsx,filename).is_file()


def test_font_preview_format_roundtrip(window):
    from adv_editor.assets import register_fonts
    from adv_editor.richtext import FONT_FILE
    filename,family=import_font(FONT,window.db.path)
    window.data['fonts'][filename]=family
    register_fonts(window.data,window.db.path)
    e=editor(window);source='{font=adv_generated/fonts/'+filename+'}{b}Bold{/b} regular{/font}'
    e.set_source(source)
    assert e.source()==source
    assert e.toPlainText()=='Bold regular'
    cursor=e.textCursor();cursor.setPosition(2);e.setTextCursor(cursor)
    assert cursor.charFormat().property(FONT_FILE).endswith(filename)
    assert cursor.charFormat().fontFamilies()[0]==family


def test_new_project_location_and_cancel(window,tmp_path,monkeypatch):
    old_data=copy.deepcopy(window.data);old_path=window.db.path
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a,**k:('',''))
    window.new_project()
    assert window.db.path==old_path and window.data==old_data
    target=tmp_path/'new-project.advdb'
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a,**k:(str(target),''))
    window.new_project()
    assert window.db.path==target and asset_folder(target).is_dir()
    assert len(window.rows)==1 and window.rows[0]['text']==''


def test_character_rename_collision_and_simultaneous_swap(window):
    from adv_editor.characters import CharacterDialog
    from adv_editor.character_ops import update_characters
    dialog=CharacterDialog(window.data)
    dialog.table.item(0,0).setText('mei')
    with pytest.raises(ValueError,match='重複'): dialog.values()
    dialog.table.item(1,0).setText('aki')
    chars,colors=dialog.values()
    before=[r['id'] for r in window.rows]
    update_characters(window.data,chars,colors,dialog.renames())
    assert window.rows[0]['character']=='mei' and window.rows[1]['character']=='aki'
    assert chars['mei']=='秋' and colors['mei']=='#e5b5a5'
    assert [r['id'] for r in window.rows]==before
    assert statement(window.rows[0]).startswith('mei ')
    dialog.close()


def action_data():
    data=sample()
    data['sections'][0]['call']=''
    data['sections'][0]['rows'].insert(0,dict(new_row(),kind='scene',text='bg classroom',transition='fade'))
    data['sections'][0]['rows'].append(dict(new_row(),kind='menu',text='要去哪裡？',choices=[dict(text='跟她走',target='greeting'),dict(text='留在原地',target='opening')]))
    return data


def test_stage_menu_storage_excel_export(tmp_path):
    data=action_data();validate(data)
    db=Database(tmp_path/'actions.advdb');db.save(data);assert db.load()==data;db.close()
    xlsx=tmp_path/'actions.xlsx';export_excel(data,xlsx);assert import_excel(xlsx)==data
    folder=export_rpy(data,tmp_path/'game')
    text=(folder/'dialogue.rpy').read_text()
    assert '    scene bg classroom with fade' in text
    assert '    menu:\n        "跟她走":\n            jump adv_greeting' in text
    assert '        "留在原地":\n            jump adv_opening' in text


def test_image_assets_copy_and_export(app,tmp_path):
    source=tmp_path/'bg.png';image=QImage(80,60,QImage.Format.Format_RGB32);image.fill(QColor('#c0d8e8'));image.save(str(source))
    project=tmp_path/'a.advdb'
    filename,tag=import_image(source,project)
    data=sample();data['images'][filename]=tag
    data['sections'][0]['rows'].insert(0,dict(new_row(),kind='scene',text=tag,transition='dissolve'))
    folder=export_rpy(data,tmp_path/'game',project)
    assert (folder/'images'/filename).read_bytes()==source.read_bytes()
    assert f'image {tag} = "adv_generated/images/{filename}"' in (folder/'characters.rpy').read_text()
    target=tmp_path/'b.advdb';copy_assets(data,project,target)
    assert image_file(target,filename).is_file()


def test_branch_links_follow_section_management():
    from adv_editor.section_ops import update_section,duplicate_sections,delete_sections,merge_sections
    data=action_data()
    update_section(data,'greeting','new_branch','第一章','新分支')
    assert data['sections'][0]['rows'][-1]['choices'][0]['target']=='new_branch'
    before=copy.deepcopy(data)
    with pytest.raises(ValueError): delete_sections(data,['new_branch'])
    assert data==before
    copies=duplicate_sections(data,['opening','new_branch'])
    copied_menu=data['sections'][2]['rows'][-1]
    assert copied_menu['choices'][0]['target']==copies[1]
    assert copied_menu['choices'][1]['target']==copies[0]


def test_text_only_hides_effects_without_modifying(window):
    window.data=action_data();window.refresh()
    before=copy.deepcopy(window.data)
    window.text_only_checkbox.setChecked(True)
    assert window.table.isRowHidden(0)
    assert window.table.isColumnHidden(0)
    assert not window.table.isRowHidden(1)
    assert window.data==before
    window.text_only_checkbox.setChecked(False)
    assert not window.table.isRowHidden(0)
    assert window.data==before


def test_stage_dialog_menu_and_insert(window,monkeypatch):
    from adv_editor.stage import StageDialog
    dialog=StageDialog(window.data,window.db.path)
    dialog.kind.setCurrentIndex(dialog.kind.findData('menu'))
    dialog.value.setText('你要怎麼做？')
    dialog.table.item(0,0).setText('繼續')
    dialog.table.cellWidget(0,1).setCurrentIndex(1)
    dialog.accept()
    assert dialog.result()==QDialog.DialogCode.Accepted
    assert dialog.result_row['choices'][0]['target']=='greeting'
    result=dialog.result_row
    def accept(d): d.result_row=result;return QDialog.DialogCode.Accepted
    monkeypatch.setattr(StageDialog,'exec',accept)
    window.add_stage('menu')
    assert window.rows[0]['kind']=='menu'
    window.undo()
    assert all(r.get('kind','dialogue')=='dialogue' for r in window.rows)


def test_manuscript_exports(tmp_path):
    from adv_editor.documents import export_word,export_pdf
    from docx import Document
    data=action_data()
    data['sections'][0]['rows'][1]['text']='前半{w}{b}加粗文字{/b}，{color=#cc0000}紅色{/color}。'
    word=tmp_path/'script.docx';pdf=tmp_path/'script.pdf'
    export_word(data,word);export_pdf(data,pdf)
    text='\n'.join(p.text for p in Document(word).paragraphs)
    assert '點擊後繼續' in text and '跟她走' in text and '台詞 ID' in text
    assert '{b}' not in text
    assert any(run.bold for p in Document(word).paragraphs for run in p.runs if '加粗文字' in run.text)
    assert pdf.read_bytes().startswith(b'%PDF') and pdf.stat().st_size>1000


def test_font_reset_and_raw_mode(window):
    from adv_editor.richtext import FONT_FAMILIES
    filename,family=import_font(FONT,window.db.path)
    path='adv_generated/fonts/'+filename
    FONT_FAMILIES[path]=family
    e=editor(window);e.set_source('{font='+path+'}A{/font}B');e.selectAll()
    e.wrap('font=')
    assert e.source()=='AB'
    window.mode.setCurrentIndex(1)
    e=window.editor;e.setPlainText('{font='+path+'}A{/font}');e.selectAll();e.wrap('font=')
    assert e.toPlainText()=='A'


def test_text_only_menu_keeps_choices_hides_targets(window):
    window.data=action_data();window.refresh()
    window.text_only_checkbox.setChecked(True)
    assert '→' not in window.table.item(3,2).text()
    assert '跟她走' in window.table.item(3,2).text()
    window.text_only_checkbox.setChecked(False)
    assert '→ greeting' in window.table.item(3,2).text()


def test_replacement_does_not_break_stage_identifiers(window):
    window.data=action_data();window.refresh()
    window.find_text.setText('classroom');window.replace_text.setText('房間')
    window.replace_all()
    assert window.rows[0]['text']=='bg classroom'
    validate(window.data)


def test_text_only_overrides_raw_view_then_restores(window):
    window.mode.setCurrentIndex(1)
    assert not window.visual
    window.text_only_checkbox.setChecked(True)
    assert window.visual and not window.mode.isEnabled()
    window.text_only_checkbox.setChecked(False)
    assert not window.visual and window.mode.isEnabled()
