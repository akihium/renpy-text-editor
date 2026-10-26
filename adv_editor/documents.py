"""Offline script manuscripts for Word and PDF. No engine markup is executed."""
from .i18n import tr
from .platform_fonts import document_font_family,pdf_font_candidates
import html
import re
from pathlib import Path


def text_runs(source):
    styles,stack = {},[]
    for part in re.split(r'(\{\{|\{[^{}\n]*\})',source):
        if not part: continue
        if part == '{{': yield '{',dict(styles); continue
        if part.startswith('{') and part.endswith('}'):
            inner = part[1:-1]
            name,_,value = inner.partition('=')
            if name in ('b','i','u','font') or (name=='color' and re.fullmatch(r'#[0-9a-fA-F]{6}',value)) or (name=='size' and value.isdigit()):
                stack.append((name,dict(styles)))
                styles[name] = value or True
            elif inner.startswith('/') and stack and stack[-1][0] == inner[1:]:
                _,styles = stack.pop()
            else:
                label = (tr('點擊後繼續') if inner in ('w','p') else tr('等待 ')+value+'s' if name in ('w','p') else tr('結束 ')+inner[1:] if inner.startswith('/') else inner)
                yield '（'+label+'）',dict(styles)
        else: yield part,dict(styles)


def manuscript_records(data,include_ids=True):
    for section in data['sections']:
        yield 'heading',section['chapter']+' / '+section['name'],None
        for row in section['rows']:
            kind = row.get('kind','dialogue')
            if kind == 'menu':
                yield 'speaker',tr('選項'),row
                if row['text']: yield 'dialogue',row['text'],row
                for choice in row['choices']:
                    target = next(s for s in data['sections'] if s['id']==choice['target'])
                    yield 'note','• '+choice['text']+' → '+target['name'],row
            elif kind != 'dialogue':
                yield 'note',tr('〔演出〕')+row['text']+(' / '+row.get('transition','') if row.get('transition') else ''),row
            else:
                cid = row['character']
                yield 'speaker',data['characters'].get(cid,tr('旁白'))+(' ('+cid+')' if cid else ''),row
                yield 'dialogue',row['text'],row
                if row['note']: yield 'note',tr('備註：')+row['note'],row
                if include_ids: yield 'id',tr('台詞 ID：')+row['id'],row


def export_word(data,path,title=None,include_ids=True):
    title=title or tr('視覺小說台詞稿')
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    default_family=document_font_family()
    section = document.sections[0]
    section.page_width,section.page_height = Cm(21),Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for style in document.styles:
        if style.type==1:
            style.font.name=default_family
            style.font.color.rgb=RGBColor(0,0,0)
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),default_family)
            for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style = document.styles[name]
        style.font.name = default_family
        style.font.color.rgb = RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),default_family)
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(7)
    title_paragraph=document.add_paragraph(title,'Title')
    for border in list(title_paragraph._p.iter(qn('w:pBdr'))): border.getparent().remove(border)
    document.add_paragraph(tr('依章節與段落整理的台詞稿，包含角色、台詞、備註及流程提示，供閱讀與校對使用。'))
    for kind,value,row in manuscript_records(data,include_ids):
        if kind=='heading': document.add_heading(value,1); continue
        paragraph = document.add_paragraph()
        if kind=='dialogue':
            for text,styles in text_runs(value):
                run = paragraph.add_run(text)
                run.bold,run.italic,run.underline = bool(styles.get('b')),bool(styles.get('i')),bool(styles.get('u'))
                if styles.get('size'): run.font.size = Pt(min(72,max(6,int(styles['size']))))
                if styles.get('color'): run.font.color.rgb = RGBColor.from_string(styles['color'][1:])
                if styles.get('font'):
                    filename = styles['font'].split('/')[-1]
                    family = data.get('fonts',{}).get(filename,default_family)
                    run.font.name = family
                    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),family)
        else:
            run = paragraph.add_run(value)
            if kind=='speaker':
                run.bold=True
                paragraph.paragraph_format.keep_with_next=True
            if kind in ('id','note'):
                run.font.size=Pt(8 if kind=='id' else 10)
                run.font.color.rgb=RGBColor.from_string('666666')
                paragraph.paragraph_format.space_after=Pt(9)
    footer=section.footer.paragraphs[0]
    footer.alignment=2
    footer.add_run(tr('頁 '))
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    document.save(path)


def export_pdf(data,path,project_path=None,title=None,include_ids=True):
    title=title or tr('視覺小說台詞稿')
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
    from .assets import font_file
    family=None
    for candidate in pdf_font_candidates():
        if not candidate.is_file(): continue
        try:
            if 'ADVBody' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont('ADVBody',str(candidate),subfontIndex=0))
            family='ADVBody'
            break
        except Exception: continue
    if family is None:
        if 'MSung-Light' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(UnicodeCIDFont('MSung-Light'))
        family='MSung-Light'
    bold_family=family
    for candidate in pdf_font_candidates(bold=True):
        if not candidate.is_file(): continue
        try:
            if 'ADVBold' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont('ADVBold',str(candidate),subfontIndex=0))
            bold_family='ADVBold'
            break
        except Exception: continue
    pdfmetrics.registerFontFamily(family,normal=family,bold=bold_family,italic=family,boldItalic=bold_family)
    fonts={}
    substitutions=[]
    for filename in data.get('fonts',{}):
        if not project_path: continue
        try:
            alias='ADV_'+Path(filename).stem
            if alias not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont(alias,str(font_file(project_path,filename))))
            pdfmetrics.registerFontFamily(alias,normal=alias,bold=alias,italic=alias,boldItalic=alias)
            fonts[filename]=alias
        except Exception:
            # ReportLab does not support all OpenType outline formats; retain readable body text.
            substitutions.append(data['fonts'][filename])
    base=ParagraphStyle('body',fontName=family,fontSize=11,leading=18,wordWrap='CJK',spaceAfter=7,autoLeading='max')
    styles={
        'title':ParagraphStyle('title',parent=base,fontName=bold_family,fontSize=21,leading=28,spaceAfter=14),
        'heading':ParagraphStyle('heading',parent=base,fontName=bold_family,fontSize=15,leading=22,spaceBefore=16,spaceAfter=10,keepWithNext=True),
        'speaker':ParagraphStyle('speaker',parent=base,fontName=bold_family,fontSize=11,leading=17,spaceBefore=7,keepWithNext=True),
        'note':ParagraphStyle('note',parent=base,fontSize=9,leading=14,textColor=colors.HexColor('#666666')),
        'id':ParagraphStyle('id',parent=base,fontSize=7,leading=11,textColor=colors.HexColor('#777777'),spaceAfter=10)}
    story=[Paragraph(html.escape(title),styles['title']),Paragraph(tr('依章節與段落整理的台詞稿，包含角色、台詞、備註及流程提示，供閱讀與校對使用。'),base),Spacer(1,8)]
    for kind,value,row in manuscript_records(data,include_ids):
        if kind=='dialogue':
            parts=[]
            for text,formatting in text_runs(value):
                safe=html.escape(text).replace('\n','<br/>')
                font=fonts.get(str(formatting.get('font','')).split('/')[-1],family)
                attrs=' name="'+font+'"'
                if formatting.get('color'): attrs+=' color="'+formatting['color']+'"'
                if formatting.get('size'): attrs+=' size="'+str(min(72,max(6,int(formatting['size']))))+'"'
                # Font file faces are retained; PDF never executes engine tags.
                if formatting.get('u'): safe='<u>'+safe+'</u>'
                if formatting.get('b'): safe='<b>'+safe+'</b>'
                if formatting.get('i'): safe='<i>'+safe+'</i>'
                parts.append('<font'+attrs+'>'+safe+'</font>')
            story.append(Paragraph(''.join(parts) or ' ',base))
        else: story.append(Paragraph(html.escape(value).replace('\n','<br/>'),styles[kind]))
    def page_number(canvas,doc):
        canvas.saveState();canvas.setFont(family,8);canvas.setFillColor(colors.HexColor('#777777'))
        canvas.drawRightString(A4[0]-55,30,str(doc.page));canvas.restoreState()
    doc=SimpleDocTemplate(str(path),pagesize=A4,leftMargin=55,rightMargin=55,topMargin=45,bottomMargin=45,title=title,author='RenPy Text Editor')
    doc.build(story,onFirstPage=page_number,onLaterPages=page_number)
    return substitutions
