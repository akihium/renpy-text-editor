from __future__ import annotations
from .i18n import tr
import copy
import json
import csv
import io
import re
import sqlite3
import uuid
from pathlib import Path


def uid(prefix):
    return prefix + uuid.uuid4().hex


def identifier(value):
    if not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*', value) or value in {'return', 'call', 'label', 'define', 'default', 'python', 'if', 'else', 'menu', 'jump', 'pass'}:
        raise ValueError(tr('ID 必須是英文字母或底線開頭，且只能包含英數字與底線。'))
    return value


def quote(text, literal=False):
    text = text.replace('\\', '\\\\').replace('"', '\\"').replace('\n', '\\n').replace('\r', '').replace('\t', '\\t')
    text = text.replace('[', '[[').replace('%', '\\%')
    if literal:
        text = text.replace('{', '{{')
    return '"' + text + '"'


def new_row():
    return {'id': uid('d_'), 'character': '', 'text': '', 'note': ''}


def sample():
    return {'characters': {'aki': '秋', 'mei': '美'}, 'character_colors': {'aki':'#e5b5a5','mei':'#b0c9ed'}, 'fonts': {}, 'images': {}, 'sections': [
        {'id': 'opening', 'chapter': '第一章', 'name': '相遇', 'call': 'greeting', 'rows': [
            dict(new_row(), character='aki', text='你好，{b}世界{/b}！', note='微笑'),
            dict(new_row(), character='mei', text='等一下{w=0.5}，一起走吧。')]},
        {'id': 'greeting', 'chapter': '第一章', 'name': '問候', 'call': '', 'rows': [dict(new_row(), text='午後的陽光灑落。')]}]}


class Database:
    def __init__(self, path):
        self.path = Path(path).resolve()
        self.conn = sqlite3.connect(self.path)
        self.conn.executescript('''PRAGMA foreign_keys=ON;
        CREATE TABLE IF NOT EXISTS characters(id TEXT PRIMARY KEY, name TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS images(filename TEXT PRIMARY KEY, tag TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS fonts(filename TEXT PRIMARY KEY, family TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS project_settings(key TEXT PRIMARY KEY, value TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS sections(id TEXT PRIMARY KEY, chapter TEXT, name TEXT, target TEXT, position INTEGER);
        CREATE TABLE IF NOT EXISTS dialogue(id TEXT PRIMARY KEY, section_id TEXT REFERENCES sections(id), character TEXT, text TEXT, note TEXT, position INTEGER);
        PRAGMA user_version=1;''')
        if 'color' not in {row[1] for row in self.conn.execute('PRAGMA table_info(characters)')}:
            self.conn.execute("ALTER TABLE characters ADD COLUMN color TEXT NOT NULL DEFAULT ''")
        if 'mode' not in {row[1] for row in self.conn.execute('PRAGMA table_info(characters)')}:
            self.conn.execute("ALTER TABLE characters ADD COLUMN mode TEXT NOT NULL DEFAULT 'adv'")
        if 'extra' not in {row[1] for row in self.conn.execute('PRAGMA table_info(dialogue)')}:
            self.conn.execute("ALTER TABLE dialogue ADD COLUMN extra TEXT NOT NULL DEFAULT '{}'")
        self.conn.execute('PRAGMA user_version=6')
        self.conn.commit()

    def load(self):
        chars = dict(self.conn.execute('SELECT id,name FROM characters'))
        sections = []
        for sid, chapter, name, target in self.conn.execute('SELECT id,chapter,name,target FROM sections ORDER BY position'):
            rows = []
            for record in self.conn.execute('SELECT id,character,text,note,extra FROM dialogue WHERE section_id=? ORDER BY position',(sid,)):
                row = dict(zip(('id','character','text','note'),record[:4]))
                row.update(json.loads(record[4]))
                rows.append(row)
            sections.append(dict(id=sid, chapter=chapter, name=name, call=target, rows=rows))
        data = dict(characters=chars, images=dict(self.conn.execute('SELECT filename,tag FROM images')), fonts=dict(self.conn.execute('SELECT filename,family FROM fonts')), character_colors={cid:color for cid,color in self.conn.execute('SELECT id,color FROM characters') if color}, sections=sections)
        prefix = self.conn.execute("SELECT value FROM project_settings WHERE key='label_prefix'").fetchone()
        if prefix is not None: data['label_prefix'] = prefix[0]
        modes = dict(self.conn.execute("SELECT id,mode FROM characters WHERE mode!='adv'"))
        if modes: data['character_modes'] = modes
        narrator = self.conn.execute("SELECT value FROM project_settings WHERE key='narrator_mode'").fetchone()
        if narrator is not None: data['narrator_mode'] = narrator[0]
        return data

    def save(self, data):
        validate(data)
        with self.conn:
            for key in ('label_prefix','narrator_mode'):
                self.conn.execute('DELETE FROM project_settings WHERE key=?',(key,))
                if key in data:
                    self.conn.execute('INSERT INTO project_settings VALUES (?,?)',(key,data[key]))
            self.conn.execute('DELETE FROM dialogue')
            self.conn.execute('DELETE FROM sections')
            self.conn.execute('DELETE FROM characters')
            self.conn.execute('DELETE FROM images')
            self.conn.executemany('INSERT INTO images VALUES (?,?)',data.get('images',{}).items())
            self.conn.execute('DELETE FROM fonts')
            self.conn.executemany('INSERT INTO fonts VALUES (?,?)',data.get('fonts',{}).items())
            self.conn.executemany('INSERT INTO characters(id,name,color,mode) VALUES (?,?,?,?)', [(cid,name,data.get('character_colors',{}).get(cid,''),data.get('character_modes',{}).get(cid,'adv')) for cid,name in data['characters'].items()])
            for i, s in enumerate(data['sections']):
                self.conn.execute('INSERT INTO sections VALUES (?,?,?,?,?)', (s['id'],s['chapter'],s['name'],s['call'],i))
                self.conn.executemany('INSERT INTO dialogue VALUES (?,?,?,?,?,?,?)', [(r['id'],s['id'],r['character'],r['text'],r['note'],j,json.dumps({k:v for k,v in r.items() if k not in ('id','character','text','note')},ensure_ascii=False)) for j,r in enumerate(s['rows'])])

    def backup(self):
        folder = self.path.parent / (self.path.stem + '_backups')
        folder.mkdir(exist_ok=True)
        from datetime import datetime
        target = folder / (datetime.now().strftime('%Y%m%d_%H%M%S_%f') + '.advdb')
        from .assets import copy_assets,asset_folder
        copy_assets(self.load(),self.path,target)
        with sqlite3.connect(target) as dest:
            self.conn.backup(dest)
        for old in sorted(folder.glob('*.advdb'))[:-20]:
            old.unlink()
            if asset_folder(old).is_dir():
                import shutil
                shutil.rmtree(asset_folder(old))

    def close(self):
        self.conn.close()


def label_name(data, section_id):
    return data.get('label_prefix','adv_') + section_id


def validate(data):
    prefix = data.get('label_prefix','adv_')
    if not isinstance(prefix,str) or (prefix and not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*',prefix)):
        raise ValueError(tr('Label 前綴只能使用英文字母、數字與底線，且不可用數字開頭；可留空。'))
    modes=data.get('character_modes',{})
    if not isinstance(modes,dict) or any(cid not in data['characters'] or mode!='nvl' for cid,mode in modes.items()):
        raise ValueError(tr('角色文字模式設定無效。'))
    if data.get('narrator_mode','adv') not in ('adv','nvl'):
        raise ValueError(tr('旁白文字模式設定無效。'))
    if (modes or data.get('narrator_mode')=='nvl') and {'nvl','nvl_narrator'} & set(data['characters']):
        raise ValueError(tr('NVL 模式下，nvl 與 nvl_narrator 是保留角色代號。'))
    ids = set()
    labels = set()
    for filename,family in data.get('fonts',{}).items():
        from .assets import font_file
        font_file('project.advdb',filename)
        if not isinstance(family,str) or not family: raise ValueError(tr('字體名稱不可空白'))
    for filename,tag in data.get('images',{}).items():
        from .assets import image_file
        image_file('project.advdb',filename)
        identifier(tag)
    for c in data['characters']:
        identifier(c)
    for cid,color in data.get('character_colors',{}).items():
        if cid not in data['characters'] or not re.fullmatch(r'#[0-9a-fA-F]{6}',color):
            raise ValueError(tr('角色顏色請使用 #RRGGBB 格式'))
    for s in data['sections']:
        identifier(s['id'])
        identifier(label_name(data,s['id']))
        if s['id'] in labels:
            raise ValueError(tr('段落 ID 重複'))
        labels.add(s['id'])
        for r in s['rows']:
            identifier(r['id'])
            if r['id'] in ids:
                raise ValueError(tr('台詞 ID 重複'))
            ids.add(r['id'])
            if r['character'] and r['character'] not in data['characters']:
                raise ValueError(tr('未知角色 ID：') + r['character'])
    for s in data['sections']:
        for row in s['rows']:
            validate_action(row,labels)
        if s['call'] and s['call'] not in labels:
            raise ValueError(tr('找不到 call 目標：') + s['call'])


def validate_action(row,labels):
    kind=row.get('kind','dialogue')
    if kind not in ('dialogue','scene','show','hide','transition','menu','jump','call','return','nvl_clear'):
        raise ValueError(tr('未知資料列類型：')+kind)
    if kind in ('scene','show','hide') and not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*(?: [A-Za-z_][A-Za-z0-9_]*)*',row['text']):
        raise ValueError(tr('圖片標籤只可使用英數字、底線及空格'))
    if row.get('transition','') not in ('','dissolve','fade','None'):
        raise ValueError(tr('不支援的轉場效果'))
    if kind=='transition' and row['text'] not in ('dissolve','fade','None'):
        raise ValueError(tr('請選擇轉場效果'))
    if kind in ('jump','call') and row['text'] not in labels:
        raise ValueError(tr('找不到分支目標：')+row['text'])
    if kind=='nvl_clear' and (row['text'] or row['character']):
        raise ValueError(tr('NVL 清頁指令不應包含角色或台詞。'))
    if any(key in row and type(row[key]) is not bool for key in ('nvl_menu','nvl_clear_after')):
        raise ValueError(tr('NVL 選項設定必須為開關值。'))
    if kind=='menu':
        if not row.get('choices'): raise ValueError(tr('選單至少需要一個選項'))
        for choice in row['choices']:
            if not choice['text'].strip(): raise ValueError(tr('選項文字不可空白'))
            if choice['target'] not in labels: raise ValueError(tr('找不到選項目標：')+choice['target'])


def statement(row, include_ids=True, label_prefix='adv_'):
    kind=row.get('kind','dialogue')
    if kind in ('scene','show','hide'):
        return kind+' '+row['text']+(' with '+row['transition'] if row.get('transition') and kind!='hide' else '')
    if kind=='transition': return 'with '+row['text']
    if kind in ('jump','call'): return kind+' '+label_prefix+row['text']
    if kind=='return': return 'return'
    if kind=='nvl_clear': return 'nvl clear'
    if kind=='menu':
        lines=[]
        nvl=row.get('nvl_menu',False)
        if row['text']: lines.append(('nvl_narrator ' if nvl else '')+quote(row['text'])+(' id '+row['id'] if include_ids else ''))
        lines.append('menu (nvl=True):' if nvl else 'menu:')
        for choice in row.get('choices',[]):
            lines.append('    '+quote(choice['text'])+':')
            if nvl and row.get('nvl_clear_after',True): lines.append('        nvl clear')
            lines.append('        jump '+label_prefix+choice['target'])
        return '\n'.join(lines)
    return (row['character'] + ' ' if row['character'] else '') + quote(row['text']) + (' id ' + row['id'] if include_ids else '')


GENERATED_HEADER = '# RenPy Text Editor — Generated file; do not edit manually'


def valid_folder_name(name):
    return bool(name.strip()) and name not in ('.', '..') and not any(c in name for c in '/\\\n\r\t\0:*?"<>|') and not name.endswith(('.', ' '))


def remap_export_fonts(text, fonts, prefix):
    # Match escaped braces as tokens so literal examples keep their original text.
    mapping = {'{font=adv_generated/fonts/'+name+'}': '{font='+prefix+'fonts/'+name+'}' for name in fonts}
    return re.sub(r'\{\{|\{[^{}\n]*\}', lambda match: mapping.get(match[0], match[0]), text)


def export_rpy(data, folder, project_path=None, include_ids=True, *, subfolder='adv_generated', game_root=None):
    validate(data)
    font_sources = []
    image_sources = []
    if data.get('images'):
        from .assets import image_file
        if project_path is None: raise ValueError(tr('匯出圖片需要專案位置'))
        for filename in data['images']:
            source=image_file(project_path,filename)
            if not source.is_file(): raise ValueError(tr('找不到專案圖片：')+str(source))
            image_sources.append((filename,source))
    if data.get('fonts'):
        from .assets import font_file
        if project_path is None: raise ValueError(tr('匯出字體需要已儲存的專案位置'))
        for filename in data['fonts']:
            source = font_file(project_path,filename)
            if not source.is_file(): raise ValueError(tr('找不到專案字體：')+str(source))
            font_sources.append((filename,source))
    base = Path(folder).expanduser().resolve()
    if subfolder:
        if not valid_folder_name(subfolder): raise ValueError(tr('無效的匯出資料夾名稱'))
    folder = base / subfolder
    root = Path(game_root).resolve() if game_root else next((p for p in (base,*base.parents) if p.name == 'game'),base)
    try:
        relative = folder.relative_to(root).as_posix()
    except ValueError:
        raise ValueError(tr('匯出資料夾必須位於 Ren’Py game 資料夾內。'))
    prefix = '' if relative == '.' else relative + '/'
    for name in ('characters.rpy','dialogue.rpy'):
        existing = folder/name
        if existing.exists() and not existing.read_text(encoding='utf-8').startswith(GENERATED_HEADER+'\n'):
            raise FileExistsError(tr('為保留手寫內容，未覆寫既有檔案：')+str(existing))
    folder.mkdir(parents=True, exist_ok=True)
    chars = GENERATED_HEADER + '\n'
    for cid, name in data['characters'].items():
        color = data.get('character_colors',{}).get(cid)
        option = ', color='+quote(color) if color else ''
        if data.get('character_modes',{}).get(cid)=='nvl': option += ', kind=nvl'
        chars += f'define {cid} = Character(_({quote(name, True)}){option})\n'
    if data.get('narrator_mode')=='nvl': chars += 'define narrator = nvl_narrator\n'
    for filename,tag in data.get('images',{}).items():
        chars += 'image '+tag+' = '+quote(prefix+'images/'+filename)+'\n'
    lines = [GENERATED_HEADER, '']
    for s in data['sections']:
        lines.append('label ' + label_name(data,s['id']) + ':')
        for r in s['rows']:
            if r['note']:
                lines.extend('    # ' + n for n in r['note'].splitlines())
            exported = copy.deepcopy(r)
            if exported.get('kind','dialogue') in ('dialogue','menu'):
                exported['text'] = remap_export_fonts(exported['text'],data.get('fonts',{}),prefix)
                for choice in exported.get('choices',[]):
                    choice['text'] = remap_export_fonts(choice['text'],data.get('fonts',{}),prefix)
            lines.extend('    '+line for line in statement(exported,include_ids,data.get('label_prefix','adv_')).splitlines())
        if s['call']:
            lines.append('    call ' + label_name(data,s['call']))
        lines.extend(['    return', ''])
    for name, content in [('characters.rpy',chars), ('dialogue.rpy','\n'.join(lines))]:
        temporary = folder / (name + '.tmp')
        temporary.write_text(content, encoding='utf-8')
        temporary.replace(folder / name)
    if image_sources:
        import shutil
        (folder/'images').mkdir(exist_ok=True)
        for filename,source in image_sources:
            target=folder/'images'/filename
            if source.resolve()!=target.resolve(): shutil.copy2(source,target)
    if font_sources:
        import shutil
        (folder/'fonts').mkdir(exist_ok=True)
        for filename,source in font_sources:
            target = folder/'fonts'/filename
            if source.resolve() != target.resolve(): shutil.copy2(source,target)
    return folder


LEGACY_HEADERS = ['章節','段落ID','段落名稱','Call目標','台詞ID','角色ID','台詞','備註']
HEADERS = LEGACY_HEADERS+['資料列類型','轉場','選項JSON']


def export_excel(data, path, project_path=None):
    if data.get('fonts') or data.get('images'):
        from .assets import copy_assets
        if project_path is None: raise ValueError(tr('匯出含素材的 Excel 需要專案位置'))
        copy_assets(data,project_path,path)
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '台詞'
    ws.append(HEADERS)
    for s in data['sections']:
        for r in s['rows']:
            ws.append([s['chapter'],s['id'],s['name'],s['call'],r['id'],r['character'],r['text'],r['note'],r.get('kind','dialogue'),r.get('transition',''),json.dumps(dict(choices=r['choices'],**{k:r[k] for k in ('nvl_menu','nvl_clear_after') if k in r}),ensure_ascii=False) if 'choices' in r and any(k in r for k in ('nvl_menu','nvl_clear_after')) else json.dumps(r['choices'],ensure_ascii=False) if 'choices' in r else ''])
    ws.freeze_panes = 'A2'
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 24
        for cell in col:
            cell.data_type = 's'  # Never turn dialogue into spreadsheet formulas.
    chars = wb.create_sheet('角色')
    chars.append(['角色ID','名稱','顏色','文字模式'])
    for cid, name in data['characters'].items():
        chars.append([cid,name,data.get('character_colors',{}).get(cid,''),data.get('character_modes',{}).get(cid,'adv')])
        for cell in chars[chars.max_row]:
            cell.data_type = 's'
    fonts = wb.create_sheet('字體')
    fonts.append(['字體檔名','字族名稱'])
    for filename,family in data.get('fonts',{}).items():
        fonts.append([filename,family])
        for cell in fonts[fonts.max_row]: cell.data_type = 's'
    images=wb.create_sheet('圖片')
    images.append(['圖片檔名','RenPy標籤'])
    for filename,tag in data.get('images',{}).items(): images.append([filename,tag])
    if any(key in data for key in ('label_prefix','narrator_mode')):
        options = wb.create_sheet('專案設定')
        options.append(['key','value'])
        for key in ('label_prefix','narrator_mode'):
            if key in data: options.append([key,data[key]])
        for row in options:
            for cell in row: cell.data_type = 's'
    wb.save(path)


def import_excel(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        headers=list(next(wb['台詞'].values))
        if headers not in (HEADERS,LEGACY_HEADERS):
            raise ValueError(tr('Excel 欄位不符合範本'))
        data = dict(characters={}, character_colors={}, fonts={}, images={}, sections=[])
        if '專案設定' in wb.sheetnames:
            for key,value,*_ in list(wb['專案設定'].values)[1:]:
                if key in ('label_prefix','narrator_mode'): data[key] = '' if value is None else str(value)
        if '圖片' in wb.sheetnames:
            for filename,tag in list(wb['圖片'].values)[1:]:
                if filename: data['images'][str(filename)]=str(tag or '')
        if '字體' in wb.sheetnames:
            for filename,family in list(wb['字體'].values)[1:]:
                if filename: data['fonts'][str(filename)] = str(family or '')
        character_rows=list(wb['角色'].values)
        character_headers=list(character_rows[0])
        color_column=character_headers.index('顏色') if '顏色' in character_headers else None
        mode_column=character_headers.index('文字模式') if '文字模式' in character_headers else None
        for values in character_rows[1:]:
            cid,name = values[:2]
            color = values[color_column] if color_column is not None and len(values)>color_column else ''
            if cid:
                if str(cid) in data['characters']:
                    raise ValueError(tr('角色 ID 重複'))
                data['characters'][str(cid)] = str(name or '')
                if color: data['character_colors'][str(cid)] = str(color)
                mode=values[mode_column] if mode_column is not None and len(values)>mode_column else 'adv'
                if mode not in (None,'','adv','nvl'): raise ValueError(tr('角色文字模式設定無效。'))
                if mode=='nvl': data.setdefault('character_modes',{})[str(cid)]='nvl'
        by_id = {}
        for values in list(wb['台詞'].values)[1:]:
            chapter,sid,name,target,rid,cid,text,note = [str(v) if v is not None else '' for v in values[:8]]
            identifier(sid)
            if sid not in by_id:
                by_id[sid] = dict(id=sid,chapter=chapter,name=name,call=target,rows=[])
                data['sections'].append(by_id[sid])
            row=dict(id=rid or uid('d_'),character=cid,text=text,note=note)
            if len(values)>8:
                kind,transition,choices=values[8:11]
                if kind and kind!='dialogue': row['kind']=str(kind)
                if transition: row['transition']=str(transition)
                if choices:
                    decoded=json.loads(choices)
                    if isinstance(decoded,dict):
                        row['choices']=decoded['choices']
                        for key in ('nvl_menu','nvl_clear_after'):
                            if key in decoded: row[key]=decoded[key]
                    else: row['choices']=decoded
            by_id[sid]['rows'].append(row)
        if not data['sections']:
            raise ValueError(tr('Excel 沒有台詞'))
        validate(data)
        return data
    finally:
        wb.close()
